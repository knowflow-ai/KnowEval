"""参考解析支持 Word、PPT、Excel：转成文字出题、核对引文。没有页码，按段落切成「页」（和 Markdown 一样）。"""

import pytest

from knoweval.build.refparse import doc_format, parse_document


def _docx(path):
    import docx

    d = docx.Document()
    d.add_heading("设备巡检制度", level=1)
    d.add_paragraph("一线设备每周巡检一次，巡检记录保存三年。")
    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "设备", "巡检周期"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "空压机", "每周"
    d.save(path)


def _pptx(path):
    from pptx import Presentation

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "季度经营回顾"
    slide.placeholders[1].text = "第三季度营收同比增长 12.5%"
    p.save(path)


def _xlsx(path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报价单"
    ws.append(["产品", "单价（元）"])
    ws.append(["知识库服务年费", 128000])
    wb.save(path)


@pytest.mark.parametrize("name, make, expected", [
    ("制度.docx", _docx, ["一线设备每周巡检一次，巡检记录保存三年。", "空压机", "每周"]),
    ("回顾.pptx", _pptx, ["季度经营回顾", "第三季度营收同比增长 12.5%"]),
    ("报价.xlsx", _xlsx, ["知识库服务年费", "128000", "报价单"]),
])
def test_office_documents_become_text_pages(tmp_path, name, make, expected):
    path = tmp_path / name
    make(path)
    assert doc_format(path) == path.suffix[1:]
    pages = parse_document(path)
    assert pages and all(p.page is None for p in pages)
    text = "\n".join(p.text for p in pages)
    for snippet in expected:
        assert snippet in text


def test_long_office_documents_are_split_like_markdown(tmp_path):
    import docx

    d = docx.Document()
    for i in range(40):
        d.add_heading(f"第{i}节", level=2)
        d.add_paragraph("制度正文。" * 80)
    path = tmp_path / "长制度.docx"
    d.save(path)
    pages = parse_document(path)
    assert len(pages) > 1 and all(len(p.text) < 4000 for p in pages)


def test_legacy_office_formats_are_explained(tmp_path):
    path = tmp_path / "旧格式.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ValueError, match="另存为"):
        doc_format(path)


def test_build_accepts_office_documents(tmp_path):
    """评测台出题：Word、PPT、Excel 不再被跳过。"""
    from knoweval.ui.builds import PARSEABLE
    assert {"docx", "pptx", "xlsx"} <= PARSEABLE


def test_corrupt_or_encrypted_office_files_are_rejected_not_read_as_garbage(tmp_path):
    """MarkItDown 解析不了时会退回当纯文本读，把乱码当正文：先认格式，不对就明确报错。"""
    path = tmp_path / "坏.docx"
    path.write_bytes(b"x")
    with pytest.raises(ValueError, match="读不出文字"):
        parse_document(path)
    wrong = tmp_path / "其实是表格.docx"
    _xlsx(wrong)        # zip 包对了，但里面不是 Word
    with pytest.raises(ValueError, match="读不出文字"):
        parse_document(wrong)
